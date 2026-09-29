"""Live API acceptance with synthetic data, isolated users and real worker subprocesses.

Run inside the disposable acceptance container. Pass configuration JSON containing
prompt_optimizer/editable_studio (or legacy optimizer connection JSON) on stdin;
never run this script against a customer database or print secrets.
"""
import base64
from contextlib import ExitStack
import json
import os
from pathlib import Path
import sys
import tempfile
import time
from unittest.mock import patch

os.environ.setdefault("CHATGPT2API_AUTH_KEY", "isolated-studio-acceptance")
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pptx import Presentation
from psd_tools import PSDImage
import numpy as np
from services import api as api_module
from services.config import ConfigStore
from services.editable_studio_config import StudioSettings
from services.editable_studio_service import StudioService
from services.user_service import UserService


def main():
    connection = json.load(sys.stdin)
    optimizer = connection.get("prompt_optimizer",connection)
    studio_settings = StudioSettings.model_validate({**connection.get("editable_studio",{}),"enabled":True,"min_account_age_seconds":0})
    root = Path("data/studio-http-acceptance").resolve()
    dense_psd = "--dense-psd" in sys.argv
    root.mkdir(parents=True, exist_ok=True)
    source = root / "synthetic-reference.png"
    picture = Image.new("RGBA", (800, 600), "#f6f5f0")
    draw = ImageDraw.Draw(picture)
    if dense_psd:
        colors = ["#d94c43","#287caf","#298a53","#b17120","#8751ad","#c14181","#277d88","#344b82","#813f2c","#8a8c23","#554f48"]
        for i,color in enumerate(colors):
            x,y = 55+(i%4)*190,55+(i//4)*180
            draw.ellipse((x,y,x+115,y+115),fill=color)
    else:
        draw.rounded_rectangle((115,180,365,460), radius=32, fill="#d94c43")
        draw.ellipse((470,240,685,455), fill="#287caf")
        draw.text((120,85), "STUDIO DEMO", fill="#172033", font_size=48)
    picture.save(source)
    image = base64.b64encode(source.read_bytes()).decode()
    results = []
    with tempfile.TemporaryDirectory(prefix="studio-http-private-") as temporary, ExitStack() as stack:
        private = Path(temporary)
        users = UserService(private / "users.db")
        config = ConfigStore(private / "config.json")
        config.update({"prompt_optimizer": optimizer, "editable_studio": studio_settings.model_dump()})
        config.path.chmod(0o600)
        users.create_user("acceptance@example.test", "synthetic-acceptance-password", quota=100)
        users.create_user("other@example.test", "synthetic-acceptance-password", quota=100)
        token, owner = users.login("acceptance@example.test", "synthetic-acceptance-password")
        other_token, _ = users.login("other@example.test", "synthetic-acceptance-password")
        studio = StudioService(config, users, root)
        for target, value in (("config", config), ("user_service", users)):
            stack.enter_context(patch.object(api_module, target, value))
        stack.enter_context(patch.object(api_module, "StudioService", return_value=studio))
        stack.enter_context(patch.object(api_module, "init_default_admin_key"))
        client = TestClient(api_module.create_app())
        stack.callback(client.close)
        studio.start()
        stack.callback(studio.stop)
        headers = {"Authorization": "Bearer " + token}
        for kind in (("psd",) if dense_psd else ("ppt", "psd")):
            brief = {"kind":kind, "page_count":3, "layer_count":12, "template_id":"business", "base64_images":[image],
                     "prompt":"制作3页产品介绍演示：封面、两种几何产品介绍、总结。红色为方形示例，蓝色为圆形示例；没有价格或业绩数据。" if kind=="ppt" else "将图中红色产品、蓝色圆形和STUDIO DEMO标题分别拆为3个前景图层，另有背景，保持位置，保留原图像素，不修补背景。"}
            if dense_psd:
                brief["prompt"] = "图中有11个独立圆形，按从上到下、从左到右顺序，每个圆形拆为独立前景层；加上背景共12层。保持原图位置和像素，不重绘、不修补背景、不要新增图层或拆分圆形。"
            response = client.post("/v1/editable-studio/plans", headers=headers, json=brief)
            if response.status_code != 202:
                raise RuntimeError("planning " + kind + ": " + response.text[:500])
            plan_task = response.json()
            deadline = time.time()+240
            while time.time()<deadline and plan_task["status"] in ("preparing","queued","running"):
                time.sleep(1)
                plan_task = client.get("/v1/editable-studio/plans/"+plan_task["id"],headers=headers).json()
            assert plan_task["status"] == "success",plan_task.get("error") or plan_task["phase"]
            draft = plan_task["result"]
            if kind == "psd":
                assert 1 <= len(draft["plan"]["layers"]) <= 11
                if dense_psd: assert len(draft["plan"]["layers"]) == 11
                assert config.get_studio_settings().max_layers == 30
            with studio.store.connect() as c:
                assert c.execute("SELECT COUNT(*) FROM user_quota_ledger WHERE reason='editable_studio_submit'").fetchone()[0] == len(results)
            print(json.dumps({"kind":kind,"stage":"plan-confirmed","count":len(draft["plan"]["slides"] or draft["plan"]["layers"])}),flush=True)
            body = {"plan_id":draft["plan_id"],"plan":draft["plan"],"expected_price":draft["price"],"client_task_id":"acceptance-"+kind}
            response = client.post("/v1/editable-studio/jobs", headers=headers, json=body)
            assert response.status_code == 200, response.text
            job_id = response.json()["id"]
            repeated = client.post("/v1/editable-studio/jobs", headers=headers, json=body)
            assert repeated.json()["id"] == job_id
            deadline = time.time()+600
            previous_phase = ""
            while time.time()<deadline:
                items = client.get("/v1/editable-studio/jobs", headers=headers).json()["items"]
                job = next(j for j in items if j["id"] == job_id)
                if job["phase"] != previous_phase:
                    print(json.dumps({"kind":kind,"stage":job["phase"]},ensure_ascii=False),flush=True)
                    previous_phase = job["phase"]
                if job["status"] in ("success", "error"):
                    break
                time.sleep(1)
            assert job["status"] == "success", job.get("error") or job["phase"]
            result = job["result"]
            file = client.get(result["primary_url"], headers=headers)
            assert file.status_code == 200 and len(file.content)>1000
            assert "no-store" in file.headers["cache-control"]
            assert client.get(result["primary_url"]).status_code == 403
            assert client.get(result["primary_url"], headers={"Authorization":"Bearer "+other_token}).status_code == 404
            folder = root/"outputs"/job_id
            if kind == "ppt":
                native = Presentation(folder/result["primary"])
                assert len(native.slides)==3
                assert all(any(shape.has_text_frame for shape in slide.shapes) for slide in native.slides)
                assert len(result["previews"])==3
            else:
                native = PSDImage.open(folder/result["primary"])
                assert len(native) == len(draft["plan"]["layers"]) + 1 <= 12
                assert np.array_equal(np.array(Image.open(folder/"preview.png").convert("RGBA")),np.array(picture))
                assert all(layer.has_pixels() and layer.topil().getbbox() for layer in native)
                assert not result["editable_text"]
            with studio.store.connect() as c:
                ledger = c.execute("SELECT COUNT(*) FROM user_quota_ledger WHERE reference=? AND reason='editable_studio_submit'",(job_id,)).fetchone()[0]
                balance = c.execute("SELECT quota FROM users WHERE id=?",(owner.id,)).fetchone()[0]
            assert ledger==1
            evidence = {"kind":kind,"status":"success","price":job["price"],"file_bytes":len(file.content),"job_id":job_id,"balance":balance,"owner_scoped":True,"deduplicated":True,"result":result}
            evidence["native_count"] = len(native.slides) if kind=="ppt" else len(native)
            results.append(evidence)
            print(json.dumps({k:v for k,v in evidence.items() if k!="result"},ensure_ascii=False),flush=True)
        (root/"report.json").write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding="utf-8")


if __name__ == "__main__":
    main()
