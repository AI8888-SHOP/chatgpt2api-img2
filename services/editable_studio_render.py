"""Fixed native renderers. Model output is validated data, never executable code."""
import argparse
import json
import os
from pathlib import Path
import subprocess
import zipfile

from PIL import Image
from services.editable_studio_models import DocumentPlan
from services.editable_studio_templates import get_template


def pptx_document(plan, images, out):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    theme = get_template(plan.template_id)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = plan.title
    prs.core_properties.author = "Editable Studio"

    def rectangle(slide,x,y,w,h,color):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(color)
        sh.line.fill.background()
        return sh

    def textbox(slide,text,x,y,w,h,size=22,color=None,bold=False):
        box = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.015)
        tf.margin_top = tf.margin_bottom = Inches(0.01)
        # Conservative CJK-aware size limit prevents overflowing narrow templates.
        units = sum(1 if ord(c)>255 else 0.53 for c in text)
        chosen = size
        while chosen>13 and units > max(1,int(w*72/chosen))*max(1,int(h*72/(chosen*1.35))):
            chosen -= 1
        if units > max(1,int(w*72/chosen))*max(1,int(h*72/(chosen*1.35))):
            raise ValueError("页面文字过多，请缩短正文后重试")
        for i,line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Noto Sans CJK SC"
            p.font.size = Pt(chosen)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color or theme["foreground"])
            p.space_after = Pt(chosen*0.35)
        return box

    for index,item in enumerate(plan.slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme["background"])
        rectangle(slide,0.65,0.52,0.52,0.065,theme["accent"])
        textbox(slide,theme["name"],0.7,6.98,10,0.22,10,theme["muted"])
        textbox(slide,f"{index+1:02d} / {len(plan.slides):02d}",11.5,6.96,1.2,0.25,11,theme["muted"])
        if item.layout in ("cover","closing"):
            textbox(slide,item.title,0.85,1.35,11.5,1.7,42,bold=True)
            rectangle(slide,0.9,3.32,1.1,0.065,theme["accent"])
            textbox(slide,"\n".join(item.body),0.9,3.75,10.7,2.45,22,theme["muted"])
        else:
            textbox(slide,item.title,0.8,0.82,11.7,1.0,30,bold=True)
            image_path = images[item.image_index] if item.image_index is not None and item.image_index<len(images) else None
            if item.layout == "chart" and item.chart_labels:
                chart_data = CategoryChartData()
                chart_data.categories = item.chart_labels
                chart_data.add_series("用户提供的数据",item.chart_values)
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(5.85),Inches(2.2),Inches(6.55),Inches(4.1),chart_data).chart
                chart.has_legend = False
                chart.chart_style = 10
                for series in chart.series:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = RGBColor.from_string(theme["accent"])
                width = 4.4
            elif image_path is not None:
                with Image.open(image_path) as im:
                    iw,ih = im.size
                scale = min(6.0/iw,4.45/ih)
                pw,ph = iw*scale,ih*scale
                slide.shapes.add_picture(str(image_path),Inches(6.15+(6-pw)/2),Inches(2.05+(4.45-ph)/2),width=Inches(pw),height=Inches(ph))
                width = 4.75
            else:
                width = 11.5
            if item.layout == "two_column" and image_path is None:
                for j,content in enumerate(item.body):
                    col,row = j%2,j//2
                    rectangle(slide,0.85+col*6.05,2.05+row*1.5,0.045,1.1,theme["accent"])
                    textbox(slide,content,1.05+col*6.05,2.08+row*1.5,5.3,1.15,21)
            else:
                step = 4.6/max(1,len(item.body))
                for j,content in enumerate(item.body):
                    rectangle(slide,0.9,2.18+j*step,0.07,0.07,theme["accent"])
                    textbox(slide,content,1.12,2.05+j*step,width-0.2,min(step-0.12,1.5),22)
        if item.notes:
            slide.notes_slide.notes_text_frame.text = item.notes
    path = out/"presentation.pptx"
    prs.save(path)
    # Reopen native elements before treating an archive with the right suffix as valid.
    reopened = Presentation(path)
    if len(reopened.slides)!=len(plan.slides) or not all(any(s.has_text_frame and s.text.strip() for s in slide.shapes) for slide in reopened.slides):
        raise ValueError("PPT 可编辑内容校验失败")
    profile = out/"lo-profile"
    subprocess.run(["libreoffice",f"-env:UserInstallation={profile.resolve().as_uri()}","--headless","--convert-to","pdf","--outdir",str(out),str(path)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=120)
    pdf = out/"presentation.pdf"
    if not pdf.is_file():
        raise ValueError("PPT 预览渲染失败")
    subprocess.run(["pdftoppm","-scale-to","1440","-png",str(pdf),str(out/"slide")],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=120)
    previews = sorted(out.glob("slide-*.png"),key=lambda p:int(p.stem.split("-")[-1]))
    if len(previews)!=len(plan.slides):
        raise ValueError("PPT 预览页数不匹配")
    with zipfile.ZipFile(out/"assets.zip","w",zipfile.ZIP_DEFLATED) as z:
        for i,image in enumerate(images):
            z.write(image,f"reference-{i+1}.png")
        z.writestr("制作方案.json",plan.model_dump_json(indent=2))
        z.write(pdf,pdf.name)
    return {"primary":"presentation.pptx","zip":"assets.zip","previews":[p.name for p in previews],"pdf":"presentation.pdf","editable_text":True,"slide_count":len(plan.slides),"warnings":plan.warnings}


def psd_document(plan, images, out):
    import cv2
    import numpy as np
    from psd_tools import PSDImage
    cv2.setNumThreads(1)
    with Image.open(images[0]) as image:
        original = image.convert("RGBA")
    w,h = original.size
    rgba = np.array(original)
    small = original.convert("RGB")
    small.thumbnail((1400,1400))
    rgb = np.array(small)
    sw,sh = small.size
    assigned = np.zeros((h,w),dtype=bool)
    extracted = []
    # Pixel ownership is not the PSD stacking order. A broad decoration box
    # can also segment every smaller subject beneath it. Reserve text/logo
    # and precise local regions first, then let broad regions use the remainder.
    # Keep the user-confirmed stack order when assembling the native PSD below.
    candidates = []
    for index,item in enumerate(plan.layers):
        x,y,bw,bh = item.box
        x0,y0 = int(x*sw/1000),int(y*sh/1000)
        x1,y1 = min(sw,int((x+bw)*sw/1000)),min(sh,int((y+bh)*sh/1000))
        if (x1-x0)*(y1-y0) > sw*sh*0.92 or min(x1-x0,y1-y0)<2:
            raise ValueError("AI 图层范围不合理，请调整图层方案")
        contour = np.zeros((sh,sw),np.uint8)
        if item.polygon:
            pts = np.array([[min(sw-1,int(px*sw/1000)),min(sh-1,int(py*sh/1000))] for px,py in item.polygon],np.int32)
            cv2.fillPoly(contour,[pts],1)
        else:
            contour[y0:y1,x0:x1]=1
        bounded = np.zeros_like(contour)
        bounded[y0:y1,x0:x1] = 1
        contour &= bounded
        candidates.append((index,item,contour))
    candidates.sort(key=lambda entry:(0 if entry[1].kind in ("text","logo") else 1,int(entry[2].sum()),-entry[0]))
    for index,item,contour in candidates:
        mask = np.full((sh,sw),cv2.GC_BGD,np.uint8)
        margin = max(3,int(min(sw,sh)*0.015))
        outer = cv2.dilate(contour,np.ones((margin*2+1,margin*2+1),np.uint8))
        mask[outer>0]=cv2.GC_PR_BGD
        mask[contour>0]=cv2.GC_PR_FGD
        ring = (outer>0)&(contour==0)
        if not ring.any() or not contour.any():
            raise ValueError("图层轮廓无效，请重新分析")
        bg = np.median(rgb[ring],axis=0)
        distance = np.sqrt(np.sum((rgb.astype(np.float32)-bg)**2,axis=2))
        values = distance[contour>0]
        seeds = (contour>0)&(distance>=max(12,float(np.percentile(values,75))))
        if seeds.sum()<4:
            raise ValueError(f"图层 {index+1} 无法可靠区分前景与背景，请调整范围或使用更清晰的原图")
        mask[seeds]=cv2.GC_FGD
        cv2.setRNGSeed(0)
        cv2.grabCut(rgb,mask,None,np.zeros((1,65),np.float64),np.zeros((1,65),np.float64),4,cv2.GC_INIT_WITH_MASK)
        foreground = (mask==cv2.GC_FGD)|(mask==cv2.GC_PR_FGD)
        # Do not let color similarity absorb adjacent unrelated elements.
        foreground &= outer>0
        full = cv2.resize(foreground.astype(np.uint8),(w,h),interpolation=cv2.INTER_NEAREST).astype(bool)
        full &= rgba[:,:,3]>0
        candidate_pixels = int(full.sum())
        full &= ~assigned
        visible_pixels = int(full.sum())
        if visible_pixels<16:
            if candidate_pixels>=16:
                raise ValueError(f"图层 {index+1} 与其他图层重复，去除重叠后没有足够的独立像素。请调整该图层范围或重新生成方案；未交付空图层")
            raise ValueError(f"图层 {index+1} 未识别到足够的可见像素，请调整范围或重新生成方案；未交付空图层")
        if visible_pixels>w*h*0.9:
            raise ValueError(f"图层 {index+1} 接近整张原图，无法作为独立元素交付，请调整拆分方案")
        assigned |= full
        pixels = rgba.copy()
        pixels[~full,3]=0
        pil = Image.fromarray(pixels)
        bounds = pil.getbbox()
        extracted.append((index,item,pil.crop(bounds),bounds))
    if not extracted:
        raise ValueError("未生成有效分层")
    background = rgba.copy()
    if plan.fill_background:
        # Deterministic local reconstruction; never claim recovery of original hidden pixels.
        repaired = cv2.inpaint(rgba[:,:,:3],assigned.astype(np.uint8)*255,5,cv2.INPAINT_TELEA)
        background[:,:,:3]=repaired
    else:
        background[assigned,3]=0
    psd = PSDImage.new("RGB",(w,h))
    base_image = Image.fromarray(background)
    base_layer = psd.create_pixel_layer(base_image,name="Background")
    base_layer.name = "背景（近似修补）" if plan.fill_background else "背景（遮挡区域透明）"
    base_image.save(out/"layer-00.png")
    manifest = [{"file":"layer-00.png","name":"背景","left":0,"top":0,"type":"pixel","background_repaired":plan.fill_background}]
    composite = base_image.copy()
    for index,(_,item,pil,bounds) in enumerate(sorted(extracted,key=lambda entry:entry[0]),1):
        name = item.name+("（像素文字，非字体层）" if item.kind=="text" else "")
        layer = psd.create_pixel_layer(pil,name=f"Layer {index}",left=bounds[0],top=bounds[1])
        layer.name = name
        filename = f"layer-{index:02d}.png"
        pil.save(out/filename)
        composite.alpha_composite(pil,dest=(bounds[0],bounds[1]))
        manifest.append({"file":filename,"name":name,"left":bounds[0],"top":bounds[1],"type":"pixel","kind":item.kind,"recognized_text":item.text})
    if not plan.fill_background and not np.array_equal(np.array(composite),rgba):
        # Invisible RGB values do not affect reconstruction; compare composited visible pixels.
        visible = rgba[:,:,3]>0
        if not np.array_equal(np.array(composite)[visible],rgba[visible]):
            raise ValueError("图层合成与原图不一致，未交付文件")
    target = out/"layers.psd"
    psd.save(target)
    reopened = PSDImage.open(target)
    if len(reopened)!=len(manifest) or reopened.size!=(w,h) or any(not layer.has_pixels() for layer in reopened):
        raise ValueError("PSD 图层结构验收失败")
    # Native layers may store alpha in a mask instead of raw pixel channels.
    # Check the visible foreground after masks, not only has_pixels()/topil().
    for index,layer in enumerate(list(reopened)[1:],1):
        rendered = layer.composite()
        if rendered is None or int((np.array(rendered.convert("RGBA"))[:,:,3]>0).sum())<16:
            raise ValueError(f"PSD 图层 {index} 的可见像素验收失败，未交付空图层")
    native_composite = np.array(reopened.composite().convert("RGBA"))
    expected_composite = np.array(composite)
    visible = expected_composite[:,:,3]>0
    if np.any(np.abs(native_composite[visible].astype(int)-expected_composite[visible].astype(int))>1):
        raise ValueError("PSD 原生图层合成验收失败")
    composite.save(out/"preview.png")
    original.save(out/"original.png")
    warnings = list(plan.warnings)+["文字为可独立移动的像素图层，不是可编辑字体层。", "原始隐藏图层、遮挡内容及字体信息无法从平面图恢复。"]
    if plan.fill_background:
        warnings.append("背景遮挡区域使用近似修补，并非原始内容；复杂纹理请人工检查。")
    (out/"manifest.json").write_text(json.dumps({"canvas":[w,h],"layers":manifest,"warnings":warnings},ensure_ascii=False,indent=2),encoding="utf-8")
    with zipfile.ZipFile(out/"assets.zip","w",zipfile.ZIP_DEFLATED) as z:
        for entry in manifest:
            z.write(out/entry["file"],entry["file"])
        z.write(out/"manifest.json","manifest.json")
        z.write(out/"preview.png","preview.png")
    return {"primary":"layers.psd","zip":"assets.zip","previews":["preview.png"],"original":"original.png","layer_count":len(manifest),"editable_text":False,"warnings":warnings}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--request",required=True)
    args = parser.parse_args()
    # Applied in the child, not preexec_fn in a multithreaded web process.
    import resource
    resource.setrlimit(resource.RLIMIT_CPU,(240,250))
    resource.setrlimit(resource.RLIMIT_FSIZE,(300*1024*1024,300*1024*1024))
    request = json.loads(Path(args.request).read_text())
    memory = max(768,min(4096,int(request.get("memory_mb",1536))))*1024*1024
    resource.setrlimit(resource.RLIMIT_AS,(memory,memory))
    plan = DocumentPlan.model_validate(request["plan"])
    out = Path(request["out"])
    images = [Path(p) for p in request["images"]]
    out.mkdir(parents=True,exist_ok=True,mode=0o700)
    try:
        result = pptx_document(plan,images,out) if plan.kind=="ppt" else psd_document(plan,images,out)
        (out/"result.json").write_text(json.dumps(result,ensure_ascii=False),encoding="utf-8")
    except Exception as exc:
        # Only fixed renderer errors are surfaced; no external response bodies or secrets.
        message = str(exc) if isinstance(exc,ValueError) else "文件生成或预览验收失败，请缩短内容后重试"
        (out/"render-error.json").write_text(json.dumps({"error":message},ensure_ascii=False),encoding="utf-8")
        raise SystemExit(1)


if __name__=="__main__":
    main()
