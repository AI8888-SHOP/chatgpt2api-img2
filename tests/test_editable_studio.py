import base64
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path
import tempfile
import time
import unittest
import os
from unittest.mock import patch

import httpx
from fastapi import HTTPException
from PIL import Image, ImageDraw

os.environ.setdefault("CHATGPT2API_AUTH_KEY", "test-studio-admin")

from services.config import ConfigStore
from services.editable_studio_config import StudioSettings, merge_studio_settings
from services.editable_studio_models import DocumentPlan, Slide, Layer, PlanRequest, GenerateRequest
from services.editable_studio_images import save_images
from services.editable_studio_provider import StudioProvider, StudioProviderError
from services.editable_studio_render import psd_document, pptx_document
from services.editable_studio_service import StudioService
from services.editable_studio_store import StudioStore
from services.prompt_optimizer_config import PromptOptimizerSettings
from services.user_service import UserService


def ppt_plan():
    return DocumentPlan(kind="ppt",title="演示",slides=[Slide(title="封面",body=["测试"],layout="cover"),Slide(title="内容",body=["可编辑正文"]),Slide(title="结束",body=["下一步"],layout="closing")])


class StudioTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.users=UserService(self.root/"users.db")
        self.user=self.users.create_user("studio@example.test","safe-test-password",quota=100)
        self.other=self.users.create_user("other@example.test","safe-test-password",quota=100)
        self.s=StudioSettings(enabled=True,reuse_optimizer_connection=False,base_url="https://provider.example/v1",api_key="secret-key-not-public",model="test-model",min_account_age_seconds=0)
        self.config=ConfigStore(self.root/"config.json")
        self.config.update({"editable_studio":self.s.model_dump()})
        self.service=StudioService(self.config,self.users,self.root/"files")
        self.store=self.service.store
        self.plan_id="a"*32
        self.store.save_plan(self.plan_id,self.user.id,self.plan_id,1,ppt_plan().model_dump(),{"prompt":"test"},[])

    def balance(self):
        with self.store.connect() as c: return c.execute("SELECT quota FROM users WHERE id=?",(self.user.id,)).fetchone()[0]

    def body(self,client="unique-client-id"):
        return GenerateRequest(plan_id=self.plan_id,client_task_id=client,plan=ppt_plan(),expected_price=3)

    def test_connection_can_reuse_disabled_optimizer_without_enabling_it(self):
        s=StudioSettings(enabled=True)
        optimizer=PromptOptimizerSettings(base_url="https://example.test/v1/chat/completions",api_key="secret",model="test")
        self.assertEqual(s.connection(optimizer),("https://example.test/v1","secret","test"))
        self.assertFalse(optimizer.enabled)

    def plan_request(self, client="plan-client-one", **changes):
        return PlanRequest(**{"kind":"ppt","prompt":"三页演示","page_count":3,"client_task_id":client,**changes})

    def install_plan_provider(self):
        calls=[]
        class FakeProvider:
            usage_tokens=150;output_tokens=50;usage_known=True
            def __init__(self,*args):pass
            def produce(self,brief,images,**kwargs):
                calls.append((brief,images,kwargs))
                return ppt_plan()
        self.service.provider_factory=FakeProvider
        return calls

    def test_async_plan_is_durable_and_deduplicated_across_threads(self):
        calls=self.install_plan_provider()
        with ThreadPoolExecutor(max_workers=5) as pool:
            tasks=list(pool.map(lambda _:self.service.plan(self.user,self.plan_request()),range(5)))
        self.assertEqual(len({t["id"] for t in tasks}),1)
        self.assertEqual(calls,[])
        restarted=StudioService(self.config,self.users,self.service.root,provider_factory=self.service.provider_factory)
        queued=restarted.store.get_plan_task(self.user.id,tasks[0]["id"])
        self.assertEqual(queued["status"],"queued")
        self.assertNotIn(self.s.api_key,queued["settings"])
        self.assertNotIn("base_url",queued["settings"])
        claimed=restarted.store.claim("worker",self.s)
        restarted.run_plan(claimed)
        result=restarted.public_plan_task(restarted.store.get_plan_task(self.user.id,tasks[0]["id"]))
        self.assertEqual(result["status"],"success")
        self.assertEqual(result["result"]["plan"]["title"],"演示")
        self.assertEqual(len(calls),1)
        self.assertEqual(calls[0][2]["deadline"],claimed["deadline"])
        self.assertEqual(self.balance(),100)
        repeated=self.service.plan(self.user,self.plan_request())
        self.assertEqual(repeated["id"],result["id"])
        with self.store.connect() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM studio_calls").fetchone()[0],1)
            self.assertEqual(c.execute("SELECT actual_tokens FROM studio_calls").fetchone()[0],150)

    def test_async_plan_duplicate_with_changed_request_rejected(self):
        self.service.plan(self.user,self.plan_request())
        with self.assertRaises(HTTPException) as error:
            self.service.plan(self.user,self.plan_request(prompt="不同需求"))
        self.assertEqual(error.exception.status_code,409)

    def test_async_plan_and_generation_share_queue_and_concurrency(self):
        self.s.global_concurrency=1;self.s.global_queue_size=1
        self.config.data["editable_studio"].update(global_concurrency=1,global_queue_size=1)
        self.install_plan_provider()
        task=self.service.plan(self.other,self.plan_request())
        self.service.submit(self.user,self.body())
        with self.assertRaises(HTTPException) as error:
            self.service.submit(self.user,self.body("second-client-id"))
        self.assertEqual(error.exception.status_code,429)
        claimed=self.store.claim("worker",self.s)
        self.assertEqual(claimed["id"],task["id"])
        self.assertIsNone(self.store.claim("worker",self.s))
        self.service.run_plan(claimed)
        self.assertIsNotNone(self.store.claim("worker",self.s))

    def test_async_plan_same_user_never_runs_with_generation(self):
        self.install_plan_provider()
        self.service.submit(self.user,self.body())
        self.service.plan(self.user,self.plan_request())
        job=self.store.claim("worker",self.s)
        self.assertNotIn("task_type",job)
        self.assertIsNone(self.store.claim("worker",self.s))
        self.store.finish(job["id"],error="test")
        self.assertEqual(self.store.claim("worker",self.s)["task_type"],"plan")

    def test_async_plan_stale_worker_and_late_result_cannot_publish(self):
        self.install_plan_provider()
        task=self.service.plan(self.user,self.plan_request())
        claimed=self.store.claim("dead-worker",self.s)
        self.store.reap();self.store.reap()
        self.service.run_plan(claimed)
        row=self.store.get_plan_task(self.user.id,task["id"])
        self.assertEqual(row["status"],"failed")
        self.assertIn("服务中断",row["error"])
        with self.assertRaises(HTTPException):self.store.get_plan(self.user.id,claimed["plan_id"])
        self.assertEqual(self.balance(),100)

    def test_async_plan_deadline_rejects_late_success(self):
        self.install_plan_provider()
        task=self.service.plan(self.user,self.plan_request())
        claimed=self.store.claim("worker",self.s)
        with self.store.connect(True) as c:c.execute("UPDATE studio_plan_tasks SET deadline=1 WHERE id=?",(task["id"],))
        self.service.run_plan(claimed)
        self.assertEqual(self.store.get_plan_task(self.user.id,task["id"])["status"],"failed")
        with self.assertRaises(HTTPException):self.store.get_plan(self.user.id,claimed["plan_id"])

    def test_async_plan_failure_keeps_budget_and_never_debits_credits(self):
        self.install_plan_provider()
        task=self.service.plan(self.user,self.plan_request())
        self.service.provider_factory.usage_known=False
        with patch.object(self.service.provider_factory,"produce",side_effect=StudioProviderError("文档 API 请求超时")):
            self.service.run_plan(self.store.claim("worker",self.s))
        row=self.store.get_plan_task(self.user.id,task["id"])
        self.assertEqual(row["status"],"failed")
        self.assertIsNone(row["actual_tokens"])
        self.assertGreater(row["budget"],0)
        self.assertEqual(self.balance(),100)

    def test_async_plan_success_is_not_overwritten_by_late_failure(self):
        self.install_plan_provider()
        task=self.service.plan(self.user,self.plan_request())
        self.service.run_plan(self.store.claim("worker",self.s))
        self.assertFalse(self.store.finish_plan_task(task["id"],error="late failure"))
        self.assertEqual(self.store.get_plan_task(self.user.id,task["id"])["status"],"success")

    def test_async_plan_preparing_timeout_is_terminal(self):
        task={"client_id":"preparing-client","fingerprint":"test","plan_id":"b"*32,"revision":1,"kind":"ppt","request":"{}","settings":"{}"}
        row=self.store.reserve_plan(self.user.id,"b"*32,100,self.s,task=task)
        with self.store.connect(True) as c:c.execute("UPDATE studio_calls SET lease=1 WHERE id=?",(row["id"],))
        self.store.reap()
        self.assertEqual(self.store.get_plan_task(self.user.id,row["id"])["status"],"failed")
        with self.assertRaises(HTTPException):self.store.queue_plan_task(row["id"],[])

    def test_async_plan_failure_cleanup_and_reference_retention(self):
        stream=BytesIO();Image.new("RGB",(40,40),"red").save(stream,"PNG")
        data=base64.b64encode(stream.getvalue()).decode()
        task=self.service.plan(self.user,self.plan_request(base64_images=[data]))
        row=self.store.get_plan_task(self.user.id,task["id"])
        folder=self.service.root/"inputs"/row["plan_id"]
        os.utime(folder,(1,1))
        self.service.prune()
        self.assertTrue(folder.is_dir())
        self.install_plan_provider()
        with patch.object(self.service.provider_factory,"produce",side_effect=StudioProviderError("test failure")):
            self.service.run_plan(self.store.claim("worker",self.s))
        self.assertFalse(folder.exists())

    def test_async_plan_limits_are_snapshotted_without_credentials(self):
        calls=self.install_plan_provider()
        task=self.service.plan(self.user,self.plan_request())
        self.config.data["editable_studio"].update(request_timeout_seconds=600)
        row=self.store.claim("worker",self.config.get_studio_settings())
        self.assertAlmostEqual(row["deadline"]-row["started"],180)
        self.service.run_plan(row)
        self.assertEqual(len(calls),1)
        public=self.service.public_plan_task(self.store.get_plan_task(self.user.id,task["id"]))
        self.assertNotIn(self.s.api_key,json.dumps(public))

    def test_settings_secret_redaction_preservation_and_clear(self):
        public=self.config.get()["editable_studio"]
        self.assertNotIn(self.s.api_key,json.dumps(public))
        self.assertTrue(public["has_api_key"])
        self.assertEqual(merge_studio_settings(public,self.s.model_dump())["api_key"],self.s.api_key)
        public.update(enabled=False,clear_api_key=True)
        self.assertEqual(merge_studio_settings(public,self.s.model_dump())["api_key"],"")

    def test_invalid_settings_do_not_leak_values_or_replace_config(self):
        with self.assertRaises(ValueError) as e: self.config.update({"editable_studio":{"api_key":"private","global_concurrency":0}})
        self.assertNotIn("private",str(e.exception))
        self.assertEqual(self.config.get_studio_settings().model,"test-model")

    def test_user_cannot_supply_model_or_tools(self):
        for extra in ({"model":"other"},{"tools":[]},{"base_url":"https://evil.test"}):
            with self.assertRaises(ValueError): PlanRequest(kind="ppt",prompt="test",**extra)

    def test_invalid_geometry_rejected(self):
        for box in ([0,0,0,100],[900,0,200,100],[-1,0,100,100]):
            with self.assertRaises(ValueError): Layer(name="test",box=box)

    def test_duplicate_submission_charges_once_across_threads(self):
        with ThreadPoolExecutor(max_workers=5) as pool:
            jobs=list(pool.map(lambda _:self.service.submit(self.user,self.body()),range(5)))
        self.assertEqual(len({j["id"] for j in jobs}),1)
        self.assertEqual(self.balance(),97)

    def test_duplicate_id_different_payload_is_rejected(self):
        self.service.submit(self.user,self.body())
        body=self.body(); body.plan.title="changed"
        with self.assertRaises(HTTPException) as e: self.service.submit(self.user,body)
        self.assertEqual(e.exception.status_code,409)
        self.assertEqual(self.balance(),97)

    def test_duplicate_receipt_survives_price_change_and_plan_expiry(self):
        original=self.service.submit(self.user,self.body())
        self.config.data["editable_studio"]["ppt_page_price"]=9
        with self.store.connect(True) as c: c.execute("UPDATE studio_plans SET created=1")
        self.assertEqual(self.service.submit(self.user,self.body())["id"],original["id"])
        self.assertEqual(self.balance(),97)

    def test_limits_snapshot_has_no_credentials(self):
        job=self.service.submit(self.user,self.body())
        snapshot=json.loads(self.store.get_job(self.user.id,job["id"])["settings"])
        self.assertEqual(snapshot["render_memory_mb"],1536)
        self.assertNotIn("api_key",snapshot)
        self.assertNotIn("base_url",snapshot)

    def test_storage_rejection_does_not_charge(self):
        self.config.data["editable_studio"]["max_storage_mb"]=256
        with self.assertRaises(HTTPException) as e: self.service.submit(self.user,self.body())
        self.assertEqual(e.exception.status_code,503)
        self.assertEqual(self.balance(),100)

    def test_refund_is_atomic_and_idempotent(self):
        j=self.service.submit(self.user,self.body())
        with ThreadPoolExecutor(max_workers=4) as pool: list(pool.map(lambda _:self.store.finish(j["id"],error="failure"),range(4)))
        self.assertEqual(self.balance(),100)
        with self.store.connect() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM user_quota_ledger WHERE reference=? AND reason='editable_studio_refund'",(j["id"],)).fetchone()[0],1)

    def test_other_user_plan_is_inaccessible(self):
        with self.assertRaises(HTTPException) as e: self.service.submit(self.other,self.body())
        self.assertEqual(e.exception.status_code,404)

    def test_price_must_be_reconfirmed(self):
        b=self.body(); b.expected_price=0
        with self.assertRaises(HTTPException) as e: self.service.submit(self.user,b)
        self.assertEqual(e.exception.status_code,409)
        self.assertEqual(self.balance(),100)

    def test_pending_budget_cannot_be_oversubscribed(self):
        self.config.data["editable_studio"]["user_daily_tokens"]=8192
        with self.assertRaises(HTTPException) as e: self.service.submit(self.user,self.body())
        self.assertEqual(e.exception.status_code,429)
        self.assertEqual(self.balance(),100)

    def test_user_concurrency_and_queue_persist(self):
        self.config.data["editable_studio"].update(generation_user_rpm=20,user_daily_tokens=10000000)
        for n in range(3): self.service.submit(self.user,self.body("client-id-"+str(n)))
        with self.assertRaises(HTTPException): self.service.submit(self.user,self.body("client-id-4"))
        other_store=StudioStore(self.users)
        self.assertIsNotNone(other_store.claim("worker",self.s))
        self.assertIsNone(other_store.claim("worker",self.s))

    def test_stale_worker_refund_is_exactly_once(self):
        self.service.submit(self.user,self.body())
        self.store.claim("dead-worker",self.s)
        self.store.reap(); self.store.reap()
        self.assertEqual(self.balance(),100)

    def test_planning_and_generation_share_global_concurrency(self):
        self.s.global_concurrency=1
        job=self.service.submit(self.user,self.body())
        self.assertIsNotNone(self.store.claim("worker",self.s))
        with self.assertRaises(HTTPException) as e: self.store.reserve_plan(self.other.id,"another-root",100,self.s)
        self.assertEqual(e.exception.status_code,429)
        self.store.finish(job["id"],error="test completed")
        call=self.store.reserve_plan(self.other.id,"another-root",100,self.s)
        self.service.submit(self.user,self.body("second-client-id"))
        self.assertIsNone(self.store.claim("worker",self.s))
        self.store.finish_plan_call(call,"success",usage=50)
        self.assertIsNotNone(self.store.claim("worker",self.s))

    def test_success_cannot_be_refunded_by_late_failure(self):
        j=self.service.submit(self.user,self.body())
        self.store.finish(j["id"],result={"primary":"x.pptx","zip":"x.zip","previews":[]},usage=500)
        self.assertFalse(self.store.finish(j["id"],error="late failure"))
        self.assertEqual(self.balance(),97)

    def test_plan_rpm_concurrency_and_revisions_count_failures(self):
        call=self.store.reserve_plan(self.user.id,"root",100,self.s)
        with self.assertRaises(HTTPException): self.store.reserve_plan(self.user.id,"root",100,self.s)
        self.store.finish_plan_call(call,"failed")
        self.s.plan_user_rpm=30;self.s.plan_revisions=1
        with self.assertRaises(HTTPException): self.store.reserve_plan(self.user.id,"root",100,self.s)

    def test_download_owner_allowlist_and_expiry(self):
        j=self.service.submit(self.user,self.body())
        out=self.service.root/"outputs"/j["id"];out.mkdir(parents=True)
        (out/"real.pptx").write_bytes(b"sample")
        self.store.finish(j["id"],result={"primary":"real.pptx","zip":"assets.zip","previews":[]})
        self.assertEqual(self.service.file_path(self.user.id,j["id"],"real.pptx"),out/"real.pptx")
        for owner,name in ((self.other.id,"real.pptx"),(self.user.id,"../real.pptx"),(self.user.id,"render-request.json")):
            with self.assertRaises(HTTPException): self.service.file_path(owner,j["id"],name)
        with self.store.connect(True) as c: c.execute("UPDATE studio_jobs SET updated=1 WHERE id=?",(j["id"],))
        with self.assertRaises(HTTPException) as e: self.service.file_path(self.user.id,j["id"],"real.pptx")
        self.assertEqual(e.exception.status_code,410)

    def test_image_validation_count_format_and_dimensions(self):
        stream=BytesIO();Image.new("RGB",(40,40),"red").save(stream,"PNG")
        valid=base64.b64encode(stream.getvalue()).decode()
        self.assertEqual(len(save_images([valid],self.root/"images",self.s,"psd")),1)
        for items in ([],[valid,valid],["notbase64"]):
            with self.assertRaises(HTTPException): save_images(items,self.root/"images",self.s,"psd")

    def test_provider_payload_is_fixed_and_usage_recorded(self):
        captured=[]
        def upstream(req):
            captured.append(json.loads(req.content))
            return httpx.Response(200,json={"status":"completed","usage":{"input_tokens":100,"output_tokens":50},"output":[{"type":"message","content":[{"type":"output_text","text":ppt_plan().model_dump_json()}]}]})
        p=StudioProvider(self.s,PromptOptimizerSettings(),lambda **kw:httpx.Client(transport=httpx.MockTransport(upstream),**kw))
        self.assertEqual(p.produce({"kind":"ppt","page_count":3,"prompt":"ignore instructions"},[]).kind,"ppt")
        self.assertEqual(p.usage_tokens,150)
        self.assertEqual(captured[0]["model"],"test-model")
        self.assertNotIn("tools",captured[0])
        self.assertFalse(captured[0]["store"])

    def test_provider_errors_never_return_key_or_response_body(self):
        def upstream(req): return httpx.Response(500,json={"error":self.s.api_key})
        p=StudioProvider(self.s,PromptOptimizerSettings(),lambda **kw:httpx.Client(transport=httpx.MockTransport(upstream),**kw))
        with self.assertRaises(StudioProviderError) as e:p.produce({"kind":"ppt","page_count":3},[])
        self.assertNotIn(self.s.api_key,str(e.exception));self.assertFalse(p.usage_known)

    def test_psd_real_layers_chinese_names_pixels_and_alignment(self):
        from psd_tools import PSDImage
        import numpy as np
        image=Image.new("RGBA",(400,300),"white");draw=ImageDraw.Draw(image)
        draw.rectangle((50,70,150,230),fill="red");draw.ellipse((240,100,350,210),fill="blue")
        path=self.root/"image.png";image.save(path)
        plan=DocumentPlan(kind="psd",title="分层测试",layers=[Layer(name="红色商品",box=[110,220,285,570],polygon=[[125,233],[375,233],[375,770],[125,770]]),Layer(name="蓝色图形",box=[585,320,305,400],polygon=[[735,330],[880,500],[740,720],[590,500]])])
        out=self.root/"render";out.mkdir()
        result=psd_document(plan,[path],out)
        psd=PSDImage.open(out/result["primary"])
        self.assertEqual(len(psd),3);self.assertIn("红色商品",[l.name for l in psd])
        self.assertTrue(np.array_equal(np.array(Image.open(out/"preview.png")),np.array(image)))
        self.assertTrue(np.array_equal(np.array(psd.composite().convert("RGBA")),np.array(image)))
        self.assertFalse(result["editable_text"])
        for index,layer in enumerate(list(psd)[1:],1):
            pixels=layer.composite().convert("RGBA")
            self.assertLess(pixels.width*pixels.height,400*300)
            self.assertTrue(layer.has_pixels())
            expected=np.array(Image.open(out/f"layer-{index:02d}.png"))[:,:,3]
            self.assertTrue(np.array_equal(np.array(pixels)[:,:,3],expected))

    def test_native_ppt_and_actual_preview(self):
        from pptx import Presentation
        out=self.root/"ppt";out.mkdir()
        result=pptx_document(ppt_plan(),[],out)
        native=Presentation(out/result["primary"])
        self.assertEqual(len(native.slides),3)
        self.assertEqual(len(result["previews"]),3)
        self.assertTrue(any(shape.has_text_frame and "可编辑正文" in shape.text for shape in native.slides[1].shapes))


if __name__=="__main__": unittest.main()
